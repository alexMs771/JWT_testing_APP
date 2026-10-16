# -*- coding: utf-8 -*-
"""Build a polished Word (.docx) POC document from POC-cf-token-manager.md.

- Manual cover page (company, title, quick-facts, clickable links)
- Static Table of Contents (always visible, no manual field update needed)
- Faithful markdown -> docx conversion: headings, paragraphs, bold/italic/inline-code,
  clickable links, bullet/number lists, blockquote callouts, code blocks, GFM tables,
  screenshots, and Mermaid diagrams (pre-rendered PNGs).
"""
import os, re
from PIL import Image
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = r"D:\TestingApp"
MD = os.path.join(ROOT, "POC-cf-token-manager.md")
OUT = os.path.join(ROOT, "cf-token-manager-POC.docx")
DIAGRAMS = os.path.join(ROOT, "docs", "diagrams")

BLUE = RGBColor(0x17, 0x40, 0x8B)
DARK = RGBColor(0x0D, 0x1B, 0x3E)
LINK = "2D6CDF"
CODEPINK = RGBColor(0xB5, 0x25, 0x7F)
GREY = RGBColor(0x5A, 0x64, 0x78)

# ---------- low-level helpers ----------

def set_shading(paragraph, fill):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), fill)
    pPr.append(shd)

def set_left_border(paragraph, color="2D6CDF", size=24):
    pPr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single'); left.set(qn('w:sz'), str(size))
    left.set(qn('w:space'), '10'); left.set(qn('w:color'), color)
    pbdr.append(left); pPr.append(pbdr)

def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), fill)
    tcPr.append(shd)

def add_hyperlink(paragraph, url, text):
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)
    run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), LINK); rPr.append(color)
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rPr.append(u)
    run.append(rPr)
    t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)

INLINE = re.compile(r'(\[[^\]]+\]\([^)]+\)|\*\*[^*]+\*\*|`[^`]+`|\*[^*\n]+\*)')

def add_inline(paragraph, text):
    text = text.replace('\\|', '|').replace('\\#', '#')
    pos = 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        tok = m.group(0)
        if tok.startswith('['):
            mm = re.match(r'\[([^\]]+)\]\(([^)]+)\)', tok)
            add_hyperlink(paragraph, mm.group(2), mm.group(1))
        elif tok.startswith('**'):
            r = paragraph.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith('`'):
            r = paragraph.add_run(tok[1:-1])
            r.font.name = 'Consolas'; r.font.size = Pt(9.5); r.font.color.rgb = CODEPINK
        elif tok.startswith('*'):
            r = paragraph.add_run(tok[1:-1]); r.italic = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])

def strip_md(text):
    text = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', text)
    return text.replace('`', '').replace('**', '').replace('\\', '')

def add_image_fit(doc, path, max_w, max_h, caption=None):
    with Image.open(path) as im:
        w, h = im.size
    ratio = w / h
    width = max_w
    if width / ratio > max_h:
        width = max_h * ratio
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(path, width=Inches(width))
    if caption:
        cp = doc.add_paragraph(); cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_after = Pt(10)
        r = cp.add_run(caption); r.italic = True; r.font.size = Pt(9); r.font.color.rgb = GREY

def add_code_block(doc, code_lines, lang=""):
    p = doc.add_paragraph()
    set_shading(p, "F3F5F9")
    pf = p.paragraph_format
    pf.space_before = Pt(6); pf.space_after = Pt(8)
    pf.left_indent = Inches(0.08); pf.right_indent = Inches(0.08)
    first = True
    for line in code_lines:
        if not first:
            p.add_run().add_break(WD_BREAK.LINE)
        first = False
        r = p.add_run(line if line else " ")
        r.font.name = 'Consolas'; r.font.size = Pt(9); r.font.color.rgb = RGBColor(0x24, 0x2A, 0x38)
    return p

def add_callout(doc, text):
    p = doc.add_paragraph()
    set_shading(p, "EEF3FF")
    set_left_border(p, "2D6CDF", 24)
    pf = p.paragraph_format
    pf.space_before = Pt(6); pf.space_after = Pt(8)
    pf.left_indent = Inches(0.12); pf.right_indent = Inches(0.08)
    add_inline(p, text)
    for r in p.runs:
        r.italic = True
        if r.font.color and r.font.color.rgb is None:
            r.font.color.rgb = RGBColor(0x33, 0x41, 0x5C)
    return p

def split_table_row(line):
    line = line.strip()
    if line.startswith('|'):
        line = line[1:]
    if line.endswith('|'):
        line = line[:-1]
    # split on unescaped pipes
    parts = re.split(r'(?<!\\)\|', line)
    return [c.strip() for c in parts]

def is_sep_row(line):
    return bool(re.match(r'^\s*\|?\s*:?-{2,}', line)) and set(line.strip()) <= set('|:- ')

def add_table(doc, rows):
    ncols = max(len(r) for r in rows)
    has_header = any(c.strip() for c in rows[0])
    table = doc.add_table(rows=0, cols=ncols)
    table.style = 'Table Grid'
    table.autofit = True
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        for ci in range(ncols):
            txt = row[ci] if ci < len(row) else ""
            cell = cells[ci]
            cell.paragraphs[0].text = ""
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(2)
            para.paragraph_format.space_before = Pt(2)
            add_inline(para, txt)
            if ri == 0 and has_header:
                shade_cell(cell, "17408B")
                for r in para.runs:
                    r.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                # header text may have been added as plain; ensure white/bold even for empty
                if not para.runs:
                    rr = para.add_run(" ")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def add_toc(doc):
    """Static, always-visible contents list built from the top-level (##) sections.
    Shows correctly in Word, LibreOffice, Google Docs and IDE .docx previews without
    needing a manual field update."""
    with open(MD, encoding='utf-8') as f:
        md_lines = f.read().split('\n')
    entries = []
    for l in md_lines:
        m = re.match(r'^##\s+(\d+)\.\s+(.*)$', l)
        if m:
            entries.append((m.group(1), strip_md(m.group(2)).strip()))
    for num, title in entries:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Inches(0.35); pf.first_line_indent = Inches(-0.35)
        pf.space_after = Pt(3)
        rn = p.add_run(f"{num}.\u2002"); rn.bold = True; rn.font.color.rgb = BLUE
        rn.font.size = Pt(11)
        rt = p.add_run(title); rt.font.color.rgb = RGBColor(0x24, 0x32, 0x4D)
        rt.font.size = Pt(11)

def add_page_number_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("cf-token-manager POC  \u00b7  Page ")
    r.font.size = Pt(8); r.font.color.rgb = GREY
    fld1 = OxmlElement('w:fldChar'); fld1.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = 'PAGE'
    fld2 = OxmlElement('w:fldChar'); fld2.set(qn('w:fldCharType'), 'end')
    run2 = p.add_run(); run2.font.size = Pt(8); run2.font.color.rgb = GREY
    run2._r.append(fld1); run2._r.append(instr); run2._r.append(fld2)

# ---------- document setup ----------

doc = Document()
normal = doc.styles['Normal']
normal.font.name = 'Calibri'; normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.12

for name, color, size in [('Heading 1', DARK, 17), ('Heading 2', BLUE, 13.5), ('Heading 3', RGBColor(0x24,0x32,0x4D), 11.5)]:
    st = doc.styles[name]
    st.font.color.rgb = color; st.font.size = Pt(size); st.font.name = 'Calibri'; st.font.bold = True

sec = doc.sections[0]
sec.top_margin = Inches(0.8); sec.bottom_margin = Inches(0.8)
sec.left_margin = Inches(0.85); sec.right_margin = Inches(0.85)

# ---------- COVER ----------

def cover_para(text, size, color=None, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=6):
    p = doc.add_paragraph(); p.alignment = align
    p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text); r.font.size = Pt(size); r.bold = bold; r.italic = italic
    if color is not None:
        r.font.color.rgb = color
    return p

cover_para("MITRAHSOFT", 15, BLUE, bold=True, before=6, after=0)
cover_para("ColdFusion Engineering \u2014 Proof of Concept", 10, GREY, after=30)

cover_para("cf-token-manager", 34, DARK, bold=True, after=2)
cover_para("A JWT Access & Refresh Token Lifecycle Package for ColdFusion (CFML)", 14, BLUE, after=14)

summary = ("A reusable ColdFusion package that gives developers a complete JWT login system "
           "\u2014 issue, verify, decode and refresh access & refresh tokens \u2014 instead of the raw "
           "encode/decode primitives that every existing package stops at. It runs on all three CF "
           "engines (Adobe ColdFusion, Lucee, BoxLang), was proven with a live ColdBox app, and is "
           "deployed publicly on the internet.")
cp = doc.add_paragraph(); cp.paragraph_format.space_after = Pt(18)
r = cp.add_run(summary); r.italic = True; r.font.size = Pt(11); r.font.color.rgb = RGBColor(0x33,0x41,0x5C)

# quick facts
facts = [
    ("Package", [("t", "cf-token-manager")]),
    ("Version", [("t", "1.0.0 (initial release)")]),
    ("Type", [("t", "ColdBox Module (service-only \u2014 no inbound routes)")]),
    ("Platform", [("t", "CommandBox + ColdBox (MVC)")]),
    ("Engine support", [("t", "Adobe ColdFusion >=2021  \u00b7  Lucee >=5.3  \u00b7  BoxLang >=1.0.0")]),
    ("Package repo", [("l", "https://github.com/alexMs771/cf-token-manager", "github.com/alexMs771/cf-token-manager")]),
    ("Install", [("c", "box install github:alexMs771/cf-token-manager")]),
    ("Live demo", [("l", "https://jwt-testing-app.onrender.com/", "https://jwt-testing-app.onrender.com/")]),
    ("Test app repo", [("l", "https://github.com/alexMs771/JWT_testing_APP", "github.com/alexMs771/JWT_testing_APP")]),
    ("Distribution", [("t", "Private / GitHub now \u2192 ForgeBox planned")]),
    ("License", [("t", "MIT")]),
    ("Status", [("t", "\u2705 POC validated on a live, deployed ColdBox application")]),
]
ft = doc.add_table(rows=0, cols=2); ft.style = 'Table Grid'
for label, segs in facts:
    cells = ft.add_row().cells
    lp = cells[0].paragraphs[0]; lp.paragraph_format.space_after = Pt(2); lp.paragraph_format.space_before = Pt(2)
    lr = lp.add_run(label); lr.bold = True; lr.font.size = Pt(10); lr.font.color.rgb = RGBColor(0xFF,0xFF,0xFF)
    shade_cell(cells[0], "17408B")
    vp = cells[1].paragraphs[0]; vp.paragraph_format.space_after = Pt(2); vp.paragraph_format.space_before = Pt(2)
    for seg in segs:
        if seg[0] == "t":
            vp.add_run(seg[1]).font.size = Pt(10)
        elif seg[0] == "c":
            rr = vp.add_run(seg[1]); rr.font.name = 'Consolas'; rr.font.size = Pt(9.5); rr.font.color.rgb = CODEPINK
        elif seg[0] == "l":
            add_hyperlink(vp, seg[1], seg[2])
ft.columns[0].width = Inches(1.7)

doc.add_paragraph().paragraph_format.space_after = Pt(10)
cover_para("Prepared by the JWT Package Team, MitrahSoft", 10, GREY, italic=True, before=10, after=0)
cover_para("Tested with the TestingApp ColdBox application \u00b7 Deployed live on Lucee 6 via CommandBox + Docker on Render", 9, GREY, italic=True, after=0)

# page break -> TOC
br = doc.add_paragraph(); br.add_run().add_break(WD_BREAK.PAGE)
h = doc.add_paragraph(); hr = h.add_run("Table of Contents"); hr.bold = True; hr.font.size = Pt(16); hr.font.color.rgb = DARK
h.paragraph_format.space_after = Pt(8)
add_toc(doc)
br2 = doc.add_paragraph(); br2.add_run().add_break(WD_BREAK.PAGE)

# ---------- parse markdown body ----------

with open(MD, encoding='utf-8') as f:
    lines = f.read().split('\n')

start = next(i for i, l in enumerate(lines) if l.startswith('## 1. Executive Summary'))
lines = lines[start:]

mermaid_idx = 0
i = 0
n = len(lines)

def flush_paragraph(buf):
    if not buf:
        return
    text = ' '.join(x.strip() for x in buf).strip()
    if text:
        p = doc.add_paragraph()
        add_inline(p, text)

while i < n:
    line = lines[i]
    stripped = line.strip()

    if not stripped:
        i += 1; continue

    # horizontal rule
    if stripped == '---':
        i += 1; continue

    # headings
    m = re.match(r'^(#{2,4})\s+(.*)$', line)
    if m:
        level = len(m.group(1))
        text = strip_md(m.group(2)).strip()
        style = {2: 'Heading 1', 3: 'Heading 2', 4: 'Heading 3'}[level]
        p = doc.add_paragraph(style=style)
        p.add_run(text)
        i += 1; continue

    # fenced code / mermaid
    if stripped.startswith('```'):
        lang = stripped[3:].strip().lower()
        j = i + 1
        code = []
        while j < n and not lines[j].strip().startswith('```'):
            code.append(lines[j]); j += 1
        if lang == 'mermaid':
            mermaid_idx += 1
            path = os.path.join(DIAGRAMS, f"d{mermaid_idx:02d}.png")
            if os.path.exists(path):
                add_image_fit(doc, path, max_w=5.7, max_h=7.2)
        else:
            add_code_block(doc, code, lang)
        i = j + 1; continue

    # blockquote
    if stripped.startswith('>'):
        buf = []
        while i < n and lines[i].strip().startswith('>'):
            buf.append(re.sub(r'^\s*>\s?', '', lines[i]))
            i += 1
        text = ' '.join(x.strip() for x in buf).strip()
        add_callout(doc, text)
        continue

    # table
    if stripped.startswith('|'):
        tlines = []
        while i < n and lines[i].strip().startswith('|'):
            tlines.append(lines[i]); i += 1
        rows = [split_table_row(l) for l in tlines if not is_sep_row(l)]
        if rows:
            add_table(doc, rows)
        continue

    # image
    im = re.match(r'^!\[([^\]]*)\]\(([^)]+)\)\s*$', stripped)
    if im:
        alt = im.group(1)
        rel = im.group(2)
        path = os.path.join(ROOT, rel.replace('/', os.sep))
        if os.path.exists(path):
            add_image_fit(doc, path, max_w=6.3, max_h=4.7, caption=alt)
        i += 1; continue

    # lists
    if re.match(r'^\s*[-*]\s+', line) or re.match(r'^\s*\d+\.\s+', line):
        while i < n:
            cur = lines[i]
            mb = re.match(r'^\s*[-*]\s+(.*)$', cur)
            mn = re.match(r'^\s*(\d+)\.\s+(.*)$', cur)
            if mb:
                item = mb.group(1)
                style = 'List Bullet'
            elif mn:
                item = mn.group(2)
                style = 'List Number'
            else:
                if cur.strip() and (cur.startswith('  ') or cur.startswith('\t')):
                    # continuation of previous item
                    prev = doc.paragraphs[-1]
                    prev.add_run(' ')
                    add_inline(prev, cur.strip())
                    i += 1; continue
                break
            p = doc.add_paragraph(style=style)
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, item)
            i += 1
        continue

    # default paragraph
    buf = []
    while i < n:
        cur = lines[i]
        s = cur.strip()
        if (not s or s == '---' or s.startswith('#') or s.startswith('```')
                or s.startswith('|') or s.startswith('>')
                or re.match(r'^\s*[-*]\s+', cur) or re.match(r'^\s*\d+\.\s+', cur)
                or re.match(r'^!\[', s)):
            break
        buf.append(cur); i += 1
    flush_paragraph(buf)

# footer with page numbers
add_page_number_footer(doc.sections[0])

doc.save(OUT)
print("SAVED", OUT, os.path.getsize(OUT), "bytes")
print("diagrams used:", mermaid_idx)
